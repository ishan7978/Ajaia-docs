import io
import json

import docx
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from .models import Document, DocumentShare
from .sanitize import clean_html

User = get_user_model()


def post_json(client, url, payload):
    return client.post(url, data=json.dumps(payload), content_type="application/json")


class BaseCase(TestCase):
    def setUp(self):
        self.alice = User.objects.create_user("alice", password="pw")
        self.bob = User.objects.create_user("bob", password="pw")
        self.carol = User.objects.create_user("carol", password="pw")
        self.doc = Document.objects.create(owner=self.alice, title="Plan", content="<p>hi</p>")

    def login(self, user):
        self.client.force_login(user)


class AccessControlTests(BaseCase):
    def test_login_required(self):
        resp = self.client.get(reverse("doc_list"))
        self.assertEqual(resp.status_code, 302)
        self.assertIn("/login/", resp["Location"])

    def test_stranger_gets_404_not_403(self):
        self.login(self.bob)
        self.assertEqual(self.client.get(reverse("doc_detail", args=[self.doc.pk])).status_code, 404)
        resp = post_json(self.client, reverse("doc_save", args=[self.doc.pk]), {"title": "x"})
        self.assertEqual(resp.status_code, 404)
        self.doc.refresh_from_db()
        self.assertEqual(self.doc.title, "Plan")

    def test_view_only_user_can_read_but_not_write(self):
        DocumentShare.objects.create(document=self.doc, user=self.bob, role="view")
        self.login(self.bob)
        self.assertEqual(self.client.get(reverse("doc_detail", args=[self.doc.pk])).status_code, 200)
        resp = post_json(self.client, reverse("doc_save", args=[self.doc.pk]), {"content": "<p>hax</p>"})
        self.assertEqual(resp.status_code, 403)
        self.doc.refresh_from_db()
        self.assertEqual(self.doc.content, "<p>hi</p>")

    def test_edit_user_can_write_but_not_share_or_delete(self):
        DocumentShare.objects.create(document=self.doc, user=self.bob, role="edit")
        self.login(self.bob)
        resp = post_json(self.client, reverse("doc_save", args=[self.doc.pk]), {"content": "<p>bob</p>"})
        self.assertEqual(resp.status_code, 200)
        self.doc.refresh_from_db()
        self.assertEqual(self.doc.content, "<p>bob</p>")
        resp = post_json(self.client, reverse("share_add", args=[self.doc.pk]), {"username": "carol"})
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(self.client.post(reverse("doc_delete", args=[self.doc.pk])).status_code, 403)
        self.assertTrue(Document.objects.filter(pk=self.doc.pk).exists())


class SharingTests(BaseCase):
    def test_owner_shares_and_recipient_sees_it_under_shared(self):
        self.login(self.alice)
        resp = post_json(
            self.client, reverse("share_add", args=[self.doc.pk]), {"username": "Bob", "role": "edit"}
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(DocumentShare.objects.get(document=self.doc, user=self.bob).role, "edit")

        self.login(self.bob)
        page = self.client.get(reverse("doc_list"))
        self.assertIn(self.doc, list(page.context["shared"]))
        self.assertNotIn(self.doc, list(page.context["owned"]))

    def test_resharing_updates_role_instead_of_duplicating(self):
        self.login(self.alice)
        url = reverse("share_add", args=[self.doc.pk])
        post_json(self.client, url, {"username": "bob", "role": "view"})
        post_json(self.client, url, {"username": "bob", "role": "edit"})
        self.assertEqual(DocumentShare.objects.filter(document=self.doc).count(), 1)
        self.assertEqual(DocumentShare.objects.get(document=self.doc).role, "edit")

    def test_share_validation(self):
        self.login(self.alice)
        url = reverse("share_add", args=[self.doc.pk])
        self.assertEqual(post_json(self.client, url, {"username": "nobody"}).status_code, 404)
        self.assertEqual(post_json(self.client, url, {"username": ""}).status_code, 400)
        self.assertEqual(post_json(self.client, url, {"username": "alice"}).status_code, 400)
        self.assertEqual(post_json(self.client, url, {"username": "bob", "role": "admin"}).status_code, 400)

    def test_unshare_revokes_access(self):
        DocumentShare.objects.create(document=self.doc, user=self.bob, role="view")
        self.login(self.alice)
        resp = post_json(self.client, reverse("share_remove", args=[self.doc.pk]), {"username": "bob"})
        self.assertEqual(resp.status_code, 200)
        self.login(self.bob)
        self.assertEqual(self.client.get(reverse("doc_detail", args=[self.doc.pk])).status_code, 404)


class DocumentCrudTests(BaseCase):
    def test_create_rename_save_reopen(self):
        self.login(self.alice)
        resp = self.client.post(reverse("doc_create"))
        doc = Document.objects.exclude(pk=self.doc.pk).get()
        self.assertRedirects(resp, reverse("doc_detail", args=[doc.pk]))

        html = "<h1>T</h1><p><strong>b</strong><em>i</em><u>u</u></p><ul><li>x</li></ul><ol><li>y</li></ol>"
        resp = post_json(self.client, reverse("doc_save", args=[doc.pk]), {"title": "Renamed", "content": html})
        self.assertEqual(resp.status_code, 200)

        page = self.client.get(reverse("doc_detail", args=[doc.pk]))
        self.assertEqual(page.context["doc"].title, "Renamed")
        self.assertEqual(page.context["doc"].content, html)  # formatting survives a round trip

    def test_save_validation(self):
        self.login(self.alice)
        url = reverse("doc_save", args=[self.doc.pk])
        self.assertEqual(post_json(self.client, url, {"title": "   "}).status_code, 400)
        self.assertEqual(post_json(self.client, url, {"title": "x" * 201}).status_code, 400)
        self.assertEqual(post_json(self.client, url, {"content": 123}).status_code, 400)
        bad = self.client.post(url, data="not json", content_type="application/json")
        self.assertEqual(bad.status_code, 400)

    def test_saved_html_is_sanitised(self):
        self.login(self.alice)
        post_json(
            self.client,
            reverse("doc_save", args=[self.doc.pk]),
            {"content": '<p onclick="x()">ok</p><script>alert(1)</script><a href="javascript:alert(1)">l</a>'},
        )
        self.doc.refresh_from_db()
        self.assertNotIn("<script", self.doc.content)
        self.assertNotIn("onclick", self.doc.content)
        self.assertNotIn("javascript:", self.doc.content)

    def test_owner_can_delete(self):
        self.login(self.alice)
        self.client.post(reverse("doc_delete", args=[self.doc.pk]))
        self.assertFalse(Document.objects.filter(pk=self.doc.pk).exists())


class UploadTests(BaseCase):
    def upload(self, name, data):
        self.login(self.alice)
        return self.client.post(
            reverse("doc_upload"), {"file": SimpleUploadedFile(name, data)}, follow=True
        )

    def test_txt_becomes_document_and_is_escaped(self):
        self.upload("notes.txt", b"First para\n\nSecond <b>para</b>")
        doc = Document.objects.get(title="notes")
        self.assertEqual(doc.owner, self.alice)
        self.assertIn("<p>First para</p>", doc.content)
        self.assertIn("&lt;b&gt;", doc.content)

    def test_markdown_is_converted(self):
        self.upload("readme.md", b"# Title\n\n- a\n- b\n\n**bold**")
        doc = Document.objects.get(title="readme")
        self.assertIn("<h1>Title</h1>", doc.content)
        self.assertIn("<li>a</li>", doc.content)
        self.assertIn("<strong>bold</strong>", doc.content)

    def test_docx_headings_formatting_and_lists(self):
        d = docx.Document()
        d.add_heading("Heading", level=1)
        p = d.add_paragraph()
        p.add_run("bold").bold = True
        p.add_run(" plain")
        d.add_paragraph("item", style="List Bullet")
        buf = io.BytesIO()
        d.save(buf)
        self.upload("report.docx", buf.getvalue())
        doc = Document.objects.get(title="report")
        self.assertIn("<h1>Heading</h1>", doc.content)
        self.assertIn("<strong>bold</strong>", doc.content)
        self.assertIn("<ul><li>item</li></ul>", doc.content)

    def test_rejects_unsupported_type_empty_and_corrupt(self):
        resp = self.upload("evil.exe", b"MZ")
        self.assertContains(resp, "Unsupported file type")
        resp = self.upload("empty.txt", b"")
        self.assertContains(resp, "empty")
        resp = self.upload("broken.docx", b"not a zip")
        self.assertContains(resp, "could not be read")
        self.assertEqual(Document.objects.count(), 1)  # only the setUp document

    def test_rejects_oversized_file(self):
        with self.settings(MAX_UPLOAD_BYTES=10):
            resp = self.upload("big.txt", b"x" * 100)
        self.assertContains(resp, "too large")


class SanitizeTests(TestCase):
    def test_keeps_toolbar_formatting(self):
        html = "<h2>a</h2><p><strong>b</strong><em>c</em><u>d</u></p><ol><li>e</li></ol>"
        self.assertEqual(clean_html(html), html)
